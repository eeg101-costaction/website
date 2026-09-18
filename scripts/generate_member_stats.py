#!/usr/bin/env python3
"""Extract summary statistics from a detailed eCOST Working Group member export.

Produces assets/data/member-stats.json with Gender, ITC, and YRI breakdowns
for display as bar charts on the EEG101 Members page.

Column names vary slightly between eCOST export versions, so the script tries
several known candidates and logs which one it used. If a column cannot be
found the chart for that category is omitted rather than failing the whole run.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any


# ---------------------------------------------------------------------------
# Known eCOST column name candidates (after field_key normalisation).
# Add new variants here if the platform changes its export headers.
# ---------------------------------------------------------------------------

GENDER_CANDIDATES = [
    "gender",
    "sex",
    "gender identity",
]

ITC_CANDIDATES = [
    "itc",
    "itc country",
    "from itc country",
    "inclusiveness target country",
    "cost itc",
    "itc member",
    "is itc",
]

YRI_CANDIDATES = [
    "young researcher",           # confirmed eCOST CA24148 export header
    "yri",
    "young researcher and innovator",
    "yri status",
    "early stage researcher",
    "esr",
    "is yri",
]

# Values treated as "Yes / positive" for ITC and YRI boolean columns.
YES_VALUES = {"y", "yes", "true", "1", "x", "✓", "checked"}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def field_key(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    return re.sub(r"\s+", " ", text.lstrip("﻿").lower())


def ascii_key(value: str) -> str:
    text = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def load_xlsx(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    """Return (headers, rows) from the first worksheet."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required.") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    raw_rows = ws.iter_rows(values_only=True)
    raw_headers = [field_key(v) for v in next(raw_rows)]
    rows = [
        {raw_headers[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(row) if i < len(raw_headers)}
        for row in raw_rows
        if any(v is not None and str(v).strip() for v in row)
    ]
    wb.close()
    return raw_headers, rows


def find_column(headers: list[str], candidates: list[str]) -> str | None:
    """Return the first candidate that appears in the header list, or None."""
    header_set = set(headers)
    for candidate in candidates:
        if candidate in header_set:
            return candidate
    return None


# ---------------------------------------------------------------------------
# Counting helpers
# ---------------------------------------------------------------------------

def count_gender(rows: list[dict[str, str]], col: str) -> list[dict[str, Any]]:
    """Count gender values, normalising to known eCOST labels."""
    raw: dict[str, int] = {}
    for row in rows:
        val = row.get(col, "").strip()
        raw[val] = raw.get(val, 0) + 1

    # Normalise to known buckets
    buckets: dict[str, int] = {
        "Male": 0,
        "Female": 0,
        "Prefer not to say": 0,
        "Other": 0,
        "Not stated": 0,
    }
    for val, count in raw.items():
        key = ascii_key(val)
        if key in {"male", "m", "man"}:
            buckets["Male"] += count
        elif key in {"female", "f", "woman"}:
            buckets["Female"] += count
        elif "prefer not" in key or "not say" in key or "undisclosed" in key:
            buckets["Prefer not to say"] += count
        elif val == "":
            buckets["Not stated"] += count
        else:
            buckets["Other"] += count

    # Drop zero-count buckets and build ordered list.
    return [{"label": k, "count": v} for k, v in buckets.items() if v > 0]


def count_boolean(rows: list[dict[str, str]], col: str, yes_label: str, no_label: str) -> list[dict[str, Any]]:
    """Count a Yes/No column into two labelled buckets."""
    yes = sum(1 for row in rows if row.get(col, "").strip().lower() in YES_VALUES)
    no = len(rows) - yes
    result = []
    if yes > 0:
        result.append({"label": yes_label, "count": yes})
    if no > 0:
        result.append({"label": no_label, "count": no})
    return result


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def generate_stats(input_path: Path, output_path: Path) -> None:
    headers, rows = load_xlsx(input_path)
    member_count = len(rows)
    print(f"Loaded {member_count} member rows from {input_path.name}.", flush=True)
    print(f"Available columns: {', '.join(headers)}", flush=True)

    charts: list[dict[str, Any]] = []
    columns_used: dict[str, str] = {}

    # Gender
    gender_col = find_column(headers, GENDER_CANDIDATES)
    if gender_col:
        columns_used["gender"] = gender_col
        print(f"Gender column: '{gender_col}'", flush=True)
        gender_data = count_gender(rows, gender_col)
        if gender_data:
            charts.append({
                "id": "gender",
                "title": "Gender split",
                "note": "Based on self-reported gender in the eCOST member profile.",
                "bars": gender_data,
            })
    else:
        print(
            f"Warning: no gender column found. Tried: {GENDER_CANDIDATES}",
            file=sys.stderr,
        )

    # ITC
    itc_col = find_column(headers, ITC_CANDIDATES)
    if itc_col:
        columns_used["itc"] = itc_col
        print(f"ITC column: '{itc_col}'", flush=True)
        itc_data = count_boolean(
            rows, itc_col,
            yes_label="ITC member",
            no_label="Non-ITC member",
        )
        if itc_data:
            charts.append({
                "id": "itc",
                "title": "Inclusiveness Target Countries (ITC)",
                "note": "ITC members are from countries designated by COST to widen participation across Europe.",
                "bars": itc_data,
            })
    else:
        print(
            f"Warning: no ITC column found. Tried: {ITC_CANDIDATES}",
            file=sys.stderr,
        )

    # YRI
    yri_col = find_column(headers, YRI_CANDIDATES)
    if yri_col:
        columns_used["yri"] = yri_col
        print(f"YRI column: '{yri_col}'", flush=True)
        yri_data = count_boolean(
            rows, yri_col,
            yes_label="Young Researcher / Innovator",
            no_label="Standard member",
        )
        if yri_data:
            charts.append({
                "id": "yri",
                "title": "Young Researcher and Innovator (YRI)",
                "note": "YRI members are early-career researchers as defined by the COST Action rules.",
                "bars": yri_data,
            })
    else:
        print(
            f"Warning: no YRI column found. Tried: {YRI_CANDIDATES}",
            file=sys.stderr,
        )

    if not charts:
        print(
            "No statistics columns were found in the export. "
            "member-stats.json will be written with an empty charts list.",
            file=sys.stderr,
        )

    stats: dict[str, Any] = {
        "generated_at": dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat(),
        "member_count": member_count,
        "columns_used": columns_used,
        "charts": charts,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(stats, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {len(charts)} chart(s) to {output_path}.", flush=True)


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate member summary statistics from a detailed eCOST export.")
    parser.add_argument("--input", type=Path, required=True, help="Detailed eCOST .xlsx export")
    parser.add_argument("--output", type=Path, default=Path("assets/data/member-stats.json"))
    args = parser.parse_args()
    try:
        generate_stats(args.input, args.output)
    except Exception as exc:
        print(f"Stats generation failed: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
